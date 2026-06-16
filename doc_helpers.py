# -*- coding: utf-8 -*-
"""Shared styling helpers for the AI career-transition study materials.
Import these in every month/curriculum generator so all .docx files look uniform.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x86, 0xC1)
GREY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x1E, 0x8A, 0x4C)
AMBER = RGBColor(0xB9, 0x77, 0x0E)
RED = RGBColor(0xB0, 0x30, 0x30)


def new_doc():
    doc = Document()
    s = doc.styles['Normal']
    s.font.name = 'Calibri'
    s.font.size = Pt(11)
    return doc


def title_block(doc, title, subtitle, meta):
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY
    if subtitle:
        st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = st.add_run(subtitle); r.italic = True; r.font.size = Pt(13); r.font.color.rgb = ACCENT
    if meta:
        m = doc.add_paragraph(); m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = m.add_run(meta); r.font.size = Pt(10); r.font.color.rgb = GREY
    doc.add_paragraph()


def h1(doc, text, color=NAVY):
    h = doc.add_heading(text, level=1)
    for run in h.runs: run.font.color.rgb = color
    return h


def h2(doc, text, color=ACCENT):
    h = doc.add_heading(text, level=2)
    for run in h.runs: run.font.color.rgb = color
    return h


def h3(doc, text, color=NAVY):
    h = doc.add_heading(text, level=3)
    for run in h.runs: run.font.color.rgb = color
    return h


def para(doc, text, bold=False, italic=False, color=None, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    r.font.size = Pt(size)
    return p


def lead(doc, label, text, label_color=ACCENT):
    """A paragraph that starts with a bold coloured label, then normal text."""
    p = doc.add_paragraph()
    r = p.add_run(label + ' '); r.bold = True; r.font.color.rgb = label_color
    p.add_run(text)
    return p


def bullet(doc, text, sub=False):
    p = doc.add_paragraph(style='List Bullet 2' if sub else 'List Bullet')
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.add_run(text)
    return p


def code_block(doc, code):
    p = doc.add_paragraph()
    r = p.add_run(code); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
    return p


def callout(doc, label, text, color=AMBER):
    p = doc.add_paragraph()
    r = p.add_run(f'[{label}] '); r.bold = True; r.font.color.rgb = color
    r2 = p.add_run(text); r2.italic = True; r2.font.color.rgb = GREY
    return p


def table_from_rows(doc, headers, rows, style='Light Grid Accent 1'):
    table = doc.add_table(rows=1, cols=len(headers)); table.style = style
    for i, hcell in enumerate(headers):
        table.rows[0].cells[i].text = hcell
        for p in table.rows[0].cells[i].paragraphs:
            for r in p.runs: r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return table


def quiz(doc, questions, title='Quiz'):
    """questions: list of dicts. Each: {'q': str, 'options': [..] (optional), 'answer': str, 'why': str (optional)}.
    Renders questions, then an answer key at the end."""
    h2(doc, title)
    for i, item in enumerate(questions, 1):
        p = doc.add_paragraph()
        r = p.add_run(f'Q{i}. '); r.bold = True
        p.add_run(item['q'])
        for j, opt in enumerate(item.get('options', [])):
            letter = chr(ord('A') + j)
            op = doc.add_paragraph(); op.paragraph_format.left_indent = Pt(24)
            op.add_run(f'{letter}) {opt}')
    h3(doc, f'{title} - Answer Key', color=GREEN)
    for i, item in enumerate(questions, 1):
        p = doc.add_paragraph()
        r = p.add_run(f'Q{i}: '); r.bold = True; r.font.color.rgb = GREEN
        p.add_run(str(item['answer']))
        if item.get('why'):
            p.add_run('  - ' + item['why']).italic = True


def page_break(doc):
    doc.add_page_break()
