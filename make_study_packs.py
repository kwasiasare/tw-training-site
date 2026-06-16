# -*- coding: utf-8 -*-
"""Generate a Word study pack (.docx) per module from a course's lesson decks.

Course-aware: pass a course slug or set COURSE_SLUG (default dami-ai-career).
  python make_study_packs.py kwasi-GenAI

For each module (grouped by `month`) it writes:
  course-materials/<course>/study-packs/Module_<n>_<title>.docx
containing, per lesson: title, summary, objectives, the lesson notes
(body_markdown rendered), the quiz with answers, and resources.
"""
import os, sys, json, glob, re
from docx import Document
from docx.shared import Pt, RGBColor

ROOT = os.path.dirname(os.path.abspath(__file__))
COURSE_SLUG = sys.argv[1] if len(sys.argv) > 1 else os.environ.get('COURSE_SLUG', 'dami-ai-career')
COURSE = os.path.join(ROOT, 'course-materials', COURSE_SLUG)
DECKS = os.path.join(COURSE, 'decks')
OUT = os.path.join(COURSE, 'study-packs')
os.makedirs(OUT, exist_ok=True)

NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(14)
    lines = code.split('\n')
    for i, ln in enumerate(lines):
        r = p.add_run(ln)
        r.font.name = 'Consolas'
        r.font.size = Pt(9)
        if i < len(lines) - 1:
            r.add_break()


def render_markdown(doc, md):
    in_code, buf = False, []
    for line in md.split('\n'):
        if line.strip().startswith('```'):
            if in_code:
                add_code(doc, '\n'.join(buf)); buf = []; in_code = False
            else:
                in_code = True
            continue
        if in_code:
            buf.append(line); continue
        s = line.rstrip()
        st = s.strip()
        if not st:
            continue
        if st.startswith('### '):
            doc.add_heading(st[4:], level=3)
        elif st.startswith('## '):
            doc.add_heading(st[3:], level=2)
        elif st.startswith('# '):
            doc.add_heading(st[2:], level=2)
        elif st.startswith(('- ', '* ')):
            doc.add_paragraph(st[2:], style='List Bullet')
        else:
            # strip simple inline markdown emphasis/backticks for readability
            txt = re.sub(r'[`*]', '', st)
            doc.add_paragraph(txt)
    if in_code and buf:
        add_code(doc, '\n'.join(buf))


def add_quiz(doc, quiz):
    doc.add_heading('Quiz', level=2)
    for i, q in enumerate(quiz, 1):
        doc.add_paragraph(f"{i}. {q['q']}")
        opts = q.get('options')
        if opts:
            for j, opt in enumerate(opts):
                doc.add_paragraph(f"{chr(65 + j)}. {opt}", style='List Bullet')
            ans = q['answer']
            label = f"{chr(65 + ans)}. {opts[ans]}" if isinstance(ans, int) else str(ans)
            p = doc.add_paragraph(); r = p.add_run(f"Answer: {label}"); r.bold = True
        else:
            p = doc.add_paragraph(); r = p.add_run(f"Answer: {q['answer']}"); r.bold = True
        if q.get('why'):
            doc.add_paragraph(f"Why: {q['why']}")


def main():
    decks = []
    for path in sorted(glob.glob(os.path.join(DECKS, '*.json'))):
        decks.append(json.load(open(path, encoding='utf-8')))
    if not decks:
        sys.exit(f"No decks in {DECKS}")
    months = {}
    for d in decks:
        months.setdefault(int(d['month']), []).append(d)

    course_title = COURSE_SLUG
    cj = os.path.join(COURSE, 'course.json')
    if os.path.exists(cj):
        course_title = json.load(open(cj, encoding='utf-8')).get('title', COURSE_SLUG)

    written = 0
    for month in sorted(months):
        lessons = sorted(months[month], key=lambda x: int(x['lesson']))
        mtitle = lessons[0]['monthTitle']
        doc = Document()
        h = doc.add_heading(f"{course_title}", level=0)
        sub = doc.add_heading(f"Module {month}: {mtitle}", level=1)
        doc.add_paragraph(f"Study pack — {len(lessons)} lessons. Notes, objectives, quizzes with answers, and resources.")
        for d in lessons:
            doc.add_page_break()
            doc.add_heading(f"Lesson {d['lesson']}: {d['title']}", level=1)
            doc.add_paragraph(d['summary'])
            if d.get('objectives'):
                doc.add_heading('What you will be able to do', level=2)
                for o in d['objectives']:
                    doc.add_paragraph(o, style='List Bullet')
            doc.add_heading('Notes', level=2)
            render_markdown(doc, d.get('body_markdown', ''))
            if d.get('quiz'):
                add_quiz(doc, d['quiz'])
            if d.get('resources'):
                doc.add_heading('Resources', level=2)
                for r in d['resources']:
                    label = r.get('label', '')
                    url = r.get('url', '')
                    doc.add_paragraph(f"{label} — {url}" if url else label, style='List Bullet')
        safe = re.sub(r'[^A-Za-z0-9]+', '_', mtitle).strip('_')
        fn = os.path.join(OUT, f"Module_{month}_{safe}.docx")
        doc.save(fn)
        print(f"Wrote {os.path.basename(fn)}  ({len(lessons)} lessons)")
        written += 1
    print(f"\nDone. {written} study packs in {OUT}")


if __name__ == '__main__':
    main()
