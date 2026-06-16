# -*- coding: utf-8 -*-
"""Generate Month 1 day-by-day schedule (.docx) and AI support-ticket project spec (.docx)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.dirname(os.path.abspath(__file__))
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x86, 0xC1)
GREY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x1E, 0x8A, 0x4C)


def new_doc():
    doc = Document()
    s = doc.styles['Normal']
    s.font.name = 'Calibri'
    s.font.size = Pt(11)
    return doc


def title_block(doc, title, subtitle, meta):
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY
    st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run(subtitle); r.italic = True; r.font.size = Pt(13); r.font.color.rgb = ACCENT
    m = doc.add_paragraph(); m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = m.add_run(meta); r.font.size = Pt(10); r.font.color.rgb = GREY
    doc.add_paragraph()


def heading(doc, text, level=1, color=NAVY):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = color
    return h


def para(doc, text, bold=False, italic=False, color=None, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    r.font.size = Pt(size)
    return p


def bullet(doc, text, sub=False):
    p = doc.add_paragraph(style='List Bullet 2' if sub else 'List Bullet')
    p.add_run(text)
    return p


# =================================================================== MONTH 1 ====
doc = new_doc()
title_block(doc,
    'Month 1 - Day-by-Day Study Schedule',
    'Python Foundations + Git (the make-or-break month)',
    '~30 hrs/week  |  5 study days/week  |  ~6 hrs/day  |  20 study days over 4 weeks')

heading(doc, 'How to use this schedule', 1)
para(doc, 'Each day has a focus, what to learn, and what to build. The daily rhythm: ~3-4 hrs '
          'learning, ~2 hrs hands-on practice, last 15 mins commit to GitHub and write one line '
          'on what you learned. Weekends are rest or catch-up. Consistency beats intensity - '
          'do not skip the practice to "watch more videos".')
para(doc, 'Primary resource: freeCodeCamp "Scientific Computing with Python" or the book '
          '"Python Crash Course" (Eric Matthes). Supplement with the official Python tutorial.',
          italic=True, color=GREY)

DAILY_RULE = 'Daily rule: 15 mins at the end - git add/commit/push + one sentence on what you learned.'

weeks = [
    ('Week 1 - Setup & Python Basics', [
        ('Day 1', 'Setup & first program',
         'Install Python, VS Code, set up a clean workspace. Learn print(), comments, running a script. '
         'Understand the terminal/command line basics (cd, ls/dir, running python files).',
         'Build: "hello world" + a script that prints a formatted intro about you.'),
        ('Day 2', 'Variables & data types',
         'Strings, integers, floats, booleans. String methods, f-strings, type conversion, input().',
         'Build: a tip calculator and a simple "mad libs" text generator using input().'),
        ('Day 3', 'Operators & conditionals',
         'Arithmetic/comparison/logical operators. if / elif / else. Truthiness.',
         'Build: a number-guessing checker and a grade calculator (score -> letter grade).'),
        ('Day 4', 'Lists & tuples',
         'Creating, indexing, slicing, appending, looping over lists. Tuples vs lists.',
         'Build: a to-do list manager (add/remove/show items in a list).'),
        ('Day 5', 'Loops',
         'for loops, while loops, range(), break/continue, nested loops.',
         'Build: multiplication table generator + a FizzBuzz program (classic interview warm-up).'),
    ]),
    ('Week 2 - Core Python', [
        ('Day 6', 'Dictionaries & sets',
         'Key-value pairs, dict methods, iterating dicts, sets and uniqueness.',
         'Build: a contact book (name -> phone/email) stored in a dictionary.'),
        ('Day 7', 'Functions',
         'Defining functions, parameters, return values, default args, scope.',
         'Build: refactor Week 1 programs into reusable functions.'),
        ('Day 8', 'More functions & modules',
         '*args/**kwargs, importing modules (math, random, datetime), writing your own module.',
         'Build: a dice-roll simulator and a password generator using random.'),
        ('Day 9', 'Error handling & files',
         'try/except, reading and writing text files, with open() pattern.',
         'Build: make the to-do list save/load from a .txt file so it persists.'),
        ('Day 10', 'Review & mini-project',
         'Consolidate Weeks 1-2. Re-attempt anything shaky.',
         'Build: a command-line quiz app (reads questions, scores the user). Push to GitHub with a README.'),
    ]),
    ('Week 3 - Git, OOP & Structure', [
        ('Day 11', 'Git & GitHub fundamentals',
         'What version control is. git init, add, commit, push, pull, status, log. Create a GitHub account, '
         'set up SSH or token, push your first repo.',
         'Build: put ALL prior projects into properly committed GitHub repos with READMEs.'),
        ('Day 12', 'Object-oriented programming I',
         'Classes, objects, attributes, methods, __init__.',
         'Build: a BankAccount class (deposit/withdraw/balance).'),
        ('Day 13', 'Object-oriented programming II',
         'Inheritance, methods overriding, __str__, basic encapsulation.',
         'Build: extend BankAccount with SavingsAccount (adds interest).'),
        ('Day 14', 'JSON & data formats',
         'What JSON is, json module (load/loads/dump/dumps), working with nested data.',
         'Build: rewrite the contact book to save/load as JSON.'),
        ('Day 15', 'Virtual environments & packages',
         'venv, pip install, requirements.txt, why isolation matters.',
         'Build: set up a venv, install the "requests" package, freeze requirements.'),
    ]),
    ('Week 4 - APIs, SQL Taste & Capstone', [
        ('Day 16', 'HTTP & APIs intro',
         'What an API is, HTTP methods, status codes, JSON responses. Using the requests library to GET data.',
         'Build: call a free public API (e.g. weather or jokes API) and print parsed results.'),
        ('Day 17', 'Working with APIs',
         'Query parameters, headers, API keys, handling errors and missing data.',
         'Build: a CLI weather app - user enters a city, you fetch and display the forecast.'),
        ('Day 18', 'SQL basics (taster)',
         'Tables, SELECT, WHERE, INSERT, basic joins. Use SQLite (built into Python).',
         'Build: a SQLite-backed version of the to-do list (sqlite3 module).'),
        ('Day 19', 'Capstone build day',
         'Combine skills: functions, files/JSON or SQLite, an API call, clean structure.',
         'Build: start a "personal dashboard" CLI - fetches weather + a quote + tracks tasks.'),
        ('Day 20', 'Capstone polish & review',
         'Finish the capstone, write a strong README (what it does, how to run, what you learned). '
         'Review the whole month; list weak spots to revisit in Month 2.',
         'Build: deploy nothing yet - just a clean, documented GitHub repo you are proud of.'),
    ]),
]

for wk_title, days in weeks:
    heading(doc, wk_title, 1)
    for day, focus, learn, build in days:
        h = doc.add_heading(level=2)
        r = h.add_run(f'{day}: {focus}'); r.font.color.rgb = ACCENT
        p = doc.add_paragraph(); r = p.add_run('Learn: '); r.bold = True; p.add_run(learn)
        p = doc.add_paragraph(); r = p.add_run('Practice/'); r.bold = True
        r2 = p.add_run(build.split(':', 1)[0] + ':'); r2.bold = True
        p.add_run(build.split(':', 1)[1] if ':' in build else build)

doc.add_paragraph()
para(doc, DAILY_RULE, bold=True, color=GREEN)

heading(doc, 'End-of-Month 1 checklist', 1)
for item in [
    'Comfortable writing Python with functions, loops, conditionals, dicts/lists.',
    'Can read/write files and JSON, and handle errors gracefully.',
    'Understand classes and objects (basic OOP).',
    'Can use Git/GitHub daily - 6+ repos pushed with READMEs.',
    'Have called a real API with the requests library.',
    'Have touched SQL via SQLite.',
    'A capstone CLI project on GitHub with a clean README.',
]:
    bullet(doc, item)
para(doc, 'If any box is unchecked, spend the first days of Month 2 closing the gap before moving '
          'on to LLM application skills.', italic=True, color=GREY)

doc.save(os.path.join(OUT, 'Month1_Daily_Schedule.docx'))
print('Saved Month1_Daily_Schedule.docx')

# ============================================================= PROJECT SPEC =====
doc = new_doc()
title_block(doc,
    'Project Spec: AI Support-Ticket Assistant',
    'Your flagship portfolio project - built to sell the L1 support background',
    'Starter build  |  Target: Month 3-4  |  Stack: Python + FastAPI/Streamlit + RAG')

heading(doc, '1. What it is (the pitch)', 1)
para(doc, 'An AI assistant that ingests a company knowledge base (help docs, past tickets, FAQs) and '
          'helps support agents resolve tickets faster. A user pastes or types a customer question; '
          'the app retrieves the most relevant knowledge, generates a suggested reply with citations, '
          'and lets the agent edit and send. This is the perfect first project for you because it sits '
          'exactly at the intersection of your L1 support experience and AI application development.')
para(doc, 'One-line CV bullet it earns you: "Built an LLM-powered support assistant (RAG over a '
          'knowledge base) that drafts cited ticket replies, using Python, FastAPI and a vector database."',
          italic=True, color=GREY)

heading(doc, '2. Core features (MVP)', 1)
for f in [
    'Knowledge ingestion: upload PDFs/markdown/text docs; the app chunks and embeds them into a vector store.',
    'Ask a question: agent enters a customer query in a simple UI.',
    'Retrieval: app finds the top-k most relevant knowledge chunks (semantic search).',
    'Answer generation: an LLM drafts a reply grounded ONLY in retrieved content.',
    'Citations: each answer shows which documents/sections it used (builds trust, reduces hallucination).',
    'Editable draft: the agent can edit the suggested reply before copying/sending.',
]:
    bullet(doc, f)

heading(doc, '3. Stretch features (after MVP works)', 1)
for f in [
    'Ticket categorisation: auto-tag the ticket (billing, technical, account...) using the LLM.',
    'Sentiment / urgency detection to prioritise the queue.',
    'Conversation memory: handle a multi-turn back-and-forth, not just one question.',
    '"Confidence" / no-answer guardrail: if retrieval is weak, say "escalate to human" instead of guessing.',
    'Simple analytics page: most common topics, avg suggested-reply length, etc.',
    'Feedback loop: thumbs up/down on answers, stored for evaluation.',
]:
    bullet(doc, f)

heading(doc, '4. Recommended tech stack', 1)
para(doc, 'Keep it deliberately small and modern. Every choice below is something employers ask about.')
rows = [
    ('Language', 'Python 3.11+'),
    ('LLM API', 'Anthropic (Claude) or OpenAI - via official SDK'),
    ('Embeddings', 'OpenAI text-embedding-3-small, or an open model via sentence-transformers'),
    ('Vector DB', 'Chroma (local, zero-setup) to start; mention Pinecone/pgvector as alternatives'),
    ('Orchestration', 'LangChain or LlamaIndex (optional) - or hand-roll the RAG loop to learn the internals'),
    ('Backend', 'FastAPI (REST endpoints) - great for your CV'),
    ('Frontend / UI', 'Streamlit for speed (MVP), or a simple React page later'),
    ('Storage', 'SQLite for tickets/feedback metadata'),
    ('Config', 'python-dotenv for API keys (.env, never commit secrets)'),
    ('Deployment', 'Hugging Face Spaces or Render (free tiers); Azure if you want cloud cred'),
]
table = doc.add_table(rows=1, cols=2); table.style = 'Light Grid Accent 1'
table.rows[0].cells[0].text = 'Layer'; table.rows[0].cells[1].text = 'Choice'
for c in table.rows[0].cells:
    for p in c.paragraphs:
        for r in p.runs: r.bold = True
for layer, choice in rows:
    cells = table.add_row().cells
    cells[0].text = layer; cells[1].text = choice

heading(doc, '5. How it works (architecture)', 1)
para(doc, 'Two pipelines: an offline "ingest" pipeline and an online "answer" pipeline.', bold=True)
para(doc, 'Ingest (run once per knowledge base):', bold=True, color=ACCENT)
for s in [
    'Load documents -> split into chunks (e.g. ~500 tokens with overlap).',
    'Create an embedding vector for each chunk.',
    'Store vectors + source metadata in Chroma.',
]:
    bullet(doc, s)
para(doc, 'Answer (run per question):', bold=True, color=ACCENT)
for s in [
    'Embed the customer question.',
    'Query Chroma for the top-k most similar chunks.',
    'Build a prompt: system instructions + retrieved context + the question.',
    'Send to the LLM; instruct it to answer ONLY from context and cite sources.',
    'Return the draft reply + citations to the UI.',
]:
    bullet(doc, s)

heading(doc, '6. Suggested repository structure', 1)
tree = (
    'ai-support-assistant/\n'
    '|-- README.md                # what it does, demo link, screenshots, how to run\n'
    '|-- requirements.txt         # pinned dependencies\n'
    '|-- .env.example             # shows needed keys (no real secrets)\n'
    '|-- .gitignore               # ignores .env, __pycache__, vector store\n'
    '|-- data/                    # sample knowledge-base docs\n'
    '|-- src/\n'
    '|   |-- __init__.py\n'
    '|   |-- config.py            # loads env vars / settings\n'
    '|   |-- ingest.py            # load + chunk + embed + store\n'
    '|   |-- retriever.py         # query the vector DB\n'
    '|   |-- generator.py         # build prompt + call the LLM\n'
    '|   |-- pipeline.py          # ties retriever + generator together\n'
    '|   `-- db.py                # SQLite for tickets/feedback\n'
    '|-- app/\n'
    '|   |-- streamlit_app.py     # the UI (MVP)\n'
    '|   `-- api.py               # FastAPI endpoints (stretch)\n'
    '|-- tests/\n'
    '|   `-- test_pipeline.py     # a few unit tests - shows you test your code\n'
    '`-- notebooks/\n'
    '    `-- experiments.ipynb    # prompt/eval experimentation\n'
)
p = doc.add_paragraph()
r = p.add_run(tree)
r.font.name = 'Consolas'; r.font.size = Pt(9.5)

heading(doc, '7. Build milestones', 1)
mrows = [
    ('M1', 'Hard-code a few text snippets, send question + context to the LLM, print a reply. (Prove the loop works.)'),
    ('M2', 'Add real embeddings + Chroma; retrieve top-k chunks instead of hard-coding context.'),
    ('M3', 'Add document ingestion (PDF/markdown) + chunking. Now it works on a real knowledge base.'),
    ('M4', 'Add citations and the "escalate if unsure" guardrail.'),
    ('M5', 'Wrap it in a Streamlit UI. Make it usable by a non-coder.'),
    ('M6', 'Deploy it, write the README with screenshots + a demo link, write a short blog post.'),
]
table = doc.add_table(rows=1, cols=2); table.style = 'Light Grid Accent 1'
table.rows[0].cells[0].text = 'Step'; table.rows[0].cells[1].text = 'Goal'
for c in table.rows[0].cells:
    for p in c.paragraphs:
        for r in p.runs: r.bold = True
for step, goal in mrows:
    cells = table.add_row().cells
    cells[0].text = step; cells[1].text = goal

heading(doc, '8. What to demonstrate in interviews', 1)
for item in [
    'Why RAG (vs fine-tuning) - cost, freshness, citations, no retraining.',
    'How you reduced hallucination (grounding + "answer only from context" + the escalate guardrail).',
    'Chunking trade-offs (size vs overlap vs retrieval quality).',
    'How you would evaluate answer quality (a small test set, thumbs up/down feedback).',
    'The support angle: how this would actually cut handle time on an L1 queue.',
]:
    bullet(doc, item)
para(doc, 'Pro tip: use REAL or realistic support docs (even open-source project docs) so the demo '
          'feels like a product, not a toy.', bold=True, color=GREEN)

doc.save(os.path.join(OUT, 'AI_Support_Ticket_Project_Spec.docx'))
print('Saved AI_Support_Ticket_Project_Spec.docx')
